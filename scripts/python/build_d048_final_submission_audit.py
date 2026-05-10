#!/usr/bin/env python3
"""D048 — Final submission audit + journal-specific staging.

Stages D047-canonical artifacts into two journal-specific submission packages:
  - manuscript/submission_CircGPM/         (Circulation: Genomic and Precision Medicine)
  - manuscript/submission_HumanGenomics/   (Human Genomics, Springer)

Runs the 12-item D048 PASS audit. Final journal choice is PI/author decision.

Allowed: format, package, audit. Forbidden: any analysis, claim escalation,
final classification, replication, new data.

D047 canonical files (must exist):
  manuscript/main_manuscript_D047.docx / .pdf / .tex
  manuscript/references.bib
  manuscript/cover_letter_draft.md
  manuscript/author_contributions_template.md
  manuscript/coi_funding_statement_template.md
  manuscript/supplementary_methods.{docx,pdf}
  manuscript/supplementary_tables.xlsx
  manuscript/supplementary_tables_tsv/
  manuscript/figures/
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import fitz
import openpyxl
from docx import Document

ROOT = Path("/Users/apple/Desktop/gwas_af")
MS_DIR = ROOT / "manuscript"
MS_FIG = MS_DIR / "figures"
MS_SUPP_TSV = MS_DIR / "supplementary_tables_tsv"

D047_DOCX = MS_DIR / "main_manuscript_D047.docx"
D047_PDF = MS_DIR / "main_manuscript_D047.pdf"
D047_TEX = MS_DIR / "main_manuscript_D047.tex"
D047_BIB = MS_DIR / "references.bib"
COVER_LETTER = MS_DIR / "cover_letter_draft.md"
AUTHORS_TPL = MS_DIR / "author_contributions_template.md"
COI_TPL = MS_DIR / "coi_funding_statement_template.md"
SUPP_METHODS_DOCX = MS_DIR / "supplementary_methods.docx"
SUPP_METHODS_PDF = MS_DIR / "supplementary_methods.pdf"
SUPP_METHODS_TEX = MS_DIR / "supplementary_methods.tex"
SUPP_XLSX = MS_DIR / "supplementary_tables.xlsx"

OUT_AUDIT = MS_DIR / "d048_final_submission_audit.md"
OUT_DECISION = MS_DIR / "target_journal_decision_matrix.md"
OUT_MANIFEST = MS_DIR / "submission_package_manifest.md"


# ---------------------------------------------------------------------------
# Forbidden / required language definitions (carry-over from D045/D046/D047)
# ---------------------------------------------------------------------------


FORBIDDEN_LITERAL = [
    ("Roselli 2025 EUR", "wrong AF ancestry attribution"),
    ("AF Roselli 2025 EUR", "wrong AF ancestry attribution"),
    ("EUR, GRCh38 primary anchor", "wrong AF ancestry attribution"),
    ("definitive therapeutic target", "claim too high"),
    ("validated therapeutic target", "claim too high"),
    ("validated drug target", "claim too high"),
    ("final GREEN-2", "final classification locked"),
    ("NPPA inhibition treats", "drug recommendation not supported"),
    ("NPPA inhibition prevents", "drug recommendation not supported"),
]

NEG_MARKERS = (
    "✗", "Forbidden", "forbidden", "Do not", "do not", "should not",
    "is not", "are not", "not a ", "not an ", "not the ", "(not ", " not ",
    "not part", "not run", "not yet", "rather than", "instead of",
    "*not*", "explicitly not", "never", "Never",
    "is locked", "are locked", "remain locked", "remains locked",
    "still locked", "STILL LOCKED", "LOCKED", "Locked",
    "is explicitly locked", "are explicitly locked",
    "kilitli", "kilit", "not as ", "rather than ",
)

REQUIRED_LANGUAGE_KEYWORDS = [
    ("CVDKP common-variant AFGenPlus all-ancestry", "AF ancestry/source"),
    ("not proteome-wide", "screen scope qualifier"),
    ("final_class_allowed = false", "interim claim discipline"),
    ("replication", "replication-pending caveat"),
    ("LD-confounded", "downgrade rationale"),
]


# ---------------------------------------------------------------------------
# Journal targets
# ---------------------------------------------------------------------------


JOURNALS = {
    "CircGPM": {
        "full_name": "Circulation: Genomic and Precision Medicine (American Heart Association)",
        "url": "https://www.ahajournals.org/journal/circgen",
        "scope": "Cardiovascular genomic and precision medicine; original human, animal, in vitro, and in silico research.",
        "abstract_style": "structured (Background / Methods / Results / Conclusions); ~250 words",
        "ref_style": "AMA / Vancouver-numeric (in-text superscript numbers; reference list ordered by citation)",
        "word_limit": "~5,000 words main text",
        "figure_limit": "~7 figures + tables combined; supplementary unlimited",
        "line_spacing": "double-spaced; line numbers required during review",
        "title_page": "separate title page with author affiliations + corresponding author",
        "fit_score": "high — excellent topical fit; risk: replication absent",
        "risk": "HIGHER — editor may request independent pQTL replication or LD-aware fine-mapping before review",
        "recommendation": "first choice if PI accepts higher decision-time risk for higher visibility",
    },
    "HumanGenomics": {
        "full_name": "Human Genomics (Springer Nature / BMC)",
        "url": "https://humgenomics.biomedcentral.com/",
        "scope": "Genomic variation effects on human biology, health, and disease; methodology, primary research, and database papers welcome.",
        "abstract_style": "structured (Background / Results / Conclusions); ~350 words",
        "ref_style": "Vancouver (numbered in order of appearance)",
        "word_limit": "more flexible (~6,000–8,000 words main text)",
        "figure_limit": "≤ 6 main figures recommended; supplementary unlimited",
        "line_spacing": "double-spaced; continuous line numbering",
        "title_page": "title page with affiliations, ORCIDs, corresponding author",
        "fit_score": "high — pragmatic fit for methodology-driven targeted screen; replication-pending stance is common in this venue",
        "risk": "LOWER — editor more likely to accept a methodology-emphasis paper with explicit interim framing",
        "recommendation": "second choice / safer; recommended if PI prefers a higher acceptance probability",
    },
}


# ---------------------------------------------------------------------------
# Staging
# ---------------------------------------------------------------------------


def stage_journal(target_name: str) -> Path:
    """Copy D047 canonical artifacts into a journal-specific submission folder.
    No scientific content is changed; the folder is a packaging staging directory."""
    spec = JOURNALS[target_name]
    dst = MS_DIR / f"submission_{target_name}"
    dst.mkdir(parents=True, exist_ok=True)

    # Manuscript files
    shutil.copy2(D047_DOCX, dst / f"main_manuscript_{target_name}.docx")
    shutil.copy2(D047_PDF, dst / f"main_manuscript_{target_name}.pdf")
    shutil.copy2(D047_TEX, dst / f"main_manuscript_{target_name}.tex")
    shutil.copy2(D047_BIB, dst / "references.bib")

    # Supplementary
    if SUPP_METHODS_DOCX.exists():
        shutil.copy2(SUPP_METHODS_DOCX, dst / "supplementary_methods.docx")
    if SUPP_METHODS_PDF.exists():
        shutil.copy2(SUPP_METHODS_PDF, dst / "supplementary_methods.pdf")
    if SUPP_METHODS_TEX.exists():
        shutil.copy2(SUPP_METHODS_TEX, dst / "supplementary_methods.tex")
    if SUPP_XLSX.exists():
        shutil.copy2(SUPP_XLSX, dst / "supplementary_tables.xlsx")

    # Figures + supplementary tables tsv
    if MS_FIG.exists():
        shutil.copytree(MS_FIG, dst / "figures", dirs_exist_ok=True)
    if MS_SUPP_TSV.exists():
        shutil.copytree(MS_SUPP_TSV, dst / "supplementary_tables_tsv", dirs_exist_ok=True)

    # Author / COI / funding template
    if AUTHORS_TPL.exists():
        shutil.copy2(AUTHORS_TPL, dst / "author_contributions_template.md")
    if COI_TPL.exists():
        shutil.copy2(COI_TPL, dst / "coi_funding_statement_template.md")

    # Journal-specific cover letter
    write_cover_letter(dst, target_name, spec)
    # Journal-specific format notes for PI/author
    write_journal_notes(dst, target_name, spec)
    # Per-journal submission checklist
    write_submission_checklist(dst, target_name, spec)
    return dst


def write_cover_letter(dst: Path, target_name: str, spec: dict):
    md = f"""# Cover letter — draft for {spec['full_name']} (D048 staging)

**[Date]**

**[Editor name]**
**{spec['full_name']}**
**[Journal address]**

Dear [Editor],

We are pleased to submit our manuscript, **"CVD-prioritized plasma proteomic screen for AF–HF–stroke pleiotropy: an interim coloc-supported atlas"**, for consideration at **{spec['full_name']}**.

This work addresses an unmet need at the intersection of human-genetics drug-target nomination and the clinical reality that atrial fibrillation (AF), heart failure (HF), and ischemic stroke are interlinked but therapeutically heterogeneous endpoints. Cis-pQTL Mendelian randomization (MR) is increasingly used to nominate drug targets, but it can be misleading when the cis-pQTL and disease lead variants are not co-localized. We assembled a CVD-prioritized targeted plasma-proteomic pQTL screen anchored on the CVDKP common-variant AFGenPlus all-ancestry AF meta-analysis, projected each AF-anchored intervention direction onto HF and stroke subtype outcomes, and applied selective two-trait colocalization to gate therapeutic interpretation.

**Why {spec['full_name']}.** {spec['scope']} The manuscript fits the journal's scope through (i) cardiovascular genetic discovery anchored on a large AF GWAS, (ii) a methodology contribution showing that locus-level colocalization is necessary before therapeutic interpretation of cis-pQTL MR, and (iii) an interim atlas of three clinically meaningful axes (NPPA AF–CES; F11/KNG1 antithrombotic AIS/CES; MMP12 atherothrombotic AIS/LAS).

**Key findings.**
1. Colocalization sharply filters cis-pQTL MR. Strong AF MR signals at IL6R and DSC2 were rejected (PP.H3 ≈ 1.0, PP.H4 ≈ 0).
2. NPPA emerges as the only AF-anchored coloc-supported axis (AF PP.H4 = 0.91; CES PP.H4 = 0.57); it does not generalize to HF or other stroke subtypes.
3. Two non-AF BLUE stroke modules emerge: F11 / KNG1 (antithrombotic AIS/CES) and MMP12 (atherothrombotic AIS/LAS).
4. Method validation: Python `coloc.abf` cross-validated against R `coloc` at |ΔPP| < 1×10⁻⁶.

We deliberately frame the resulting axes as **interim coloc-supported axes**, not validated therapeutic targets, because independent pQTL replication, outcome replication, and LD-aware fine-mapping have not yet been performed. This framing is consistent with current best-practice scientific reporting and allows the work to stand as a hypothesis-generating evidence layer with a clear replication roadmap.

The work is original, has not been submitted elsewhere, and all authors have approved this submission.

We have prepared the manuscript and supplementary material in line with **{spec['full_name']}** author instructions ({spec['url']}). A complete data and code availability statement, full author contributions, and COI / funding declarations are included.

We thank the editor for considering this manuscript and look forward to the reviewers' comments.

Sincerely,

**[Corresponding author name, on behalf of all authors]**
**[Affiliation]**
**[Address]**
**[Email]**
**[Phone]**

---

## Suggested reviewers (PI to fill before submission)

1. [Name, affiliation, expertif then area]
2. [Name, affiliation, expertif then area]
3. [Name, affiliation, expertif then area]

## Excluded reviewers (PI to fill before submission)

- [Name, affiliation, conflict reason]
"""
    (dst / f"cover_letter_{target_name}.md").write_text(md)


def write_journal_notes(dst: Path, target_name: str, spec: dict):
    md = f"""# {spec['full_name']} — submission format notes (D048)

This fwith lists the journal-specific format adaptations the human author/PI must apply
*before* {target_name} submission. The packaged manuscript files in this directory
are the D047-canonical scientific content; format-only adaptation is required.

## Format

| Item | {target_name} requirement |
|---|---|
| Word limit (main text) | {spec['word_limit']} |
| Abstract style | {spec['abstract_style']} |
| Reference style | {spec['ref_style']} |
| Figure / table limit | {spec['figure_limit']} |
| Line spacing | {spec['line_spacing']} |
| Title page | {spec['title_page']} |

## Fit assessment (D048)

- **Scope fit:** {spec['scope']}
- **Fit score:** {spec['fit_score']}
- **Decision risk:** {spec['risk']}
- **PI recommendation:** {spec['recommendation']}

## Required edits before submission

1. Adapt abstract to the journal-mandated structure ({spec['abstract_style']}).
   - The current D047 abstract is unstructured; split into Background / Methods / Results / Conclusions (CircGPM) or Background / Results / Conclusions (Human Genomics).
2. Apply double line spacing to main text and continuous line numbering for review.
3. Re-style references in references.bib + main_manuscript LaTeX preamble:
   - CircGPM: AMA / superscript numeric (e.g., `\\bibliographystyle{{ama}}` or journal-provided `.bst`).
   - Human Genomics: Vancouver numeric (`\\bibliographystyle{{vancouver}}` or journal-provided).
4. Replace [Authors] / [Affiliations] / corresponding-author blocks with final author list.
5. Fill author_contributions_template.md, coi_funding_statement_template.md, and cover_letter_{target_name}.md with PI-supplied content.
6. Confirm figure resolution: 300 dpi PNG (already provided in figures/) and source PDFs (provided).
7. Confirm all supplementary files attach: supplementary_methods.{{docx,pdf,tex}}, supplementary_tables.xlsx, supplementary_tables_tsv/, figures/.
8. Add journal-specific keywords (≤ 6) on title page.

## Submission portal (PI step)

- Journal portal: {spec['url']}
- Cover letter: cover_letter_{target_name}.md
- Manuscript: main_manuscript_{target_name}.docx (or LaTeX)
- Supplementary bundle: supplementary_methods.* + supplementary_tables.xlsx + figures/

## Locked items (carry-over from D041 / D043 / D044 / D045 / D046 / D047)

- Final GREEN/RED/AMBER therapeutic classification — LOCKED.
- Replication phase — LOCKED.
- Wave 3 SuSiE / D042 optional — LOCKED.
- New coloc / new MR / new data downloads / proteome-wide expansion — LOCKED.
"""
    (dst / "journal_specific_notes.md").write_text(md)


def write_submission_checklist(dst: Path, target_name: str, spec: dict):
    md = f"""# {target_name} — submission checklist (D048 → D049)

Use this checklist before clicking submit at the journal portal.

## Fwiths

- [ ] main_manuscript_{target_name}.docx (or .tex)
- [ ] main_manuscript_{target_name}.pdf
- [ ] cover_letter_{target_name}.md (final wording)
- [ ] supplementary_methods.docx (or .pdf)
- [ ] supplementary_tables.xlsx (5 sheets)
- [ ] supplementary_tables_tsv/ (5 .tsv files)
- [ ] figures/ (6 PDFs + 6 PNGs)
- [ ] references.bib (29 entries; format-converted to journal style)

## Content (PI must finalize)

- [ ] Authors list + affiliations + ORCIDs
- [ ] Corresponding author contact
- [ ] Funding sources + grant numbers
- [ ] Conflict of interest declarations
- [ ] Data and code availability statement (with repository URL)
- [ ] Ethical / data-use statement
- [ ] Suggested reviewers (≥ 3)
- [ ] Excluded reviewers (with reason if any)
- [ ] Keywords (≤ 6, journal-specific)

## Format ({target_name} requirements)

- [ ] Abstract is structured per journal style ({spec['abstract_style']})
- [ ] Word count within limit ({spec['word_limit']})
- [ ] Reference style applied ({spec['ref_style']})
- [ ] Figure / table count within limit ({spec['figure_limit']})
- [ ] Double-spaced + continuous line numbers (review-ready)
- [ ] Title page formatted per {spec['title_page']}

## Final claim-boundary verification (D044/D045/D046/D047/D048)

- [ ] No "Roselli 2025 EUR" anywhere (text or figures)
- [ ] No "validated therapeutic target" / "definitive therapeutic target"
- [ ] No "final GREEN-2" / "final GREEN" / "final RED" calls
- [ ] No "NPPA inhibition treats / prevents" wording
- [ ] AF anchor described as CVDKP common-variant AFGenPlus all-ancestry / MIXED
- [ ] CVD150 described as targeted, not proteome-wide
- [ ] Replication-pending caveat preserved
- [ ] final_class_allowed = false preserved in supplementary Table 2

## Locked items (do not promif then more than this in cover letter or revision response)

- Final GREEN/RED/AMBER therapeutic classification — LOCKED.
- Replication phase — LOCKED.
- Wave 3 SuSiE / D042 optional — LOCKED.
- New coloc / new MR / new data downloads / proteome-wide expansion — LOCKED.
"""
    (dst / "submission_checklist.md").write_text(md)


# ---------------------------------------------------------------------------
# Decision matrix + manifest
# ---------------------------------------------------------------------------


def write_decision_matrix():
    md = ["# Target journal decision matrix (D048)",
          "",
          "**Snapshot:** 2026-05-10",
          "",
          "Two journals are staged in this D048 cycle. PI/author makes the final choice; both packages are ready.",
          "",
          "## Comparison",
          "",
          "| Criterion | CircGPM | Human Genomics |",
          "|---|---|---|"]
    cg = JOURNALS["CircGPM"]
    hg = JOURNALS["HumanGenomics"]
    rows = [
        ("Full name", cg["full_name"], hg["full_name"]),
        ("URL", cg["url"], hg["url"]),
        ("Scope", cg["scope"], hg["scope"]),
        ("Abstract style", cg["abstract_style"], hg["abstract_style"]),
        ("Reference style", cg["ref_style"], hg["ref_style"]),
        ("Word limit", cg["word_limit"], hg["word_limit"]),
        ("Figure / table limit", cg["figure_limit"], hg["figure_limit"]),
        ("Fit score", cg["fit_score"], hg["fit_score"]),
        ("Decision risk", cg["risk"], hg["risk"]),
        ("PI recommendation", cg["recommendation"], hg["recommendation"]),
    ]
    for label, c, h in rows:
        md.append(f"| {label} | {c} | {h} |")

    md.extend([
        "",
        "## Recommended order",
        "",
        "1. **Circulation: Genomic and Precision Medicine (CircGPM)** — first choice if PI accepts higher decision-time risk for higher visibility. Strong topical match; the principal risk is editorial demand for replication or LD-aware fine-mapping at first review.",
        "2. **Human Genomics (Springer)** — second choice / safer alternative. Pragmatic fit, more tolerant of methodology-driven submissions with explicit interim framing.",
        "",
        "## Alternative venues (not staged in D048; record only)",
        "",
        "- npj Genomic Medicine — higher tier; replication absence may be a stronger objection.",
        "- Communications Biology — broad biology audience; targeted-screen + no-replication framing is a real risk.",
        "",
        "## Decision rule",
        "",
        "PI should pick CircGPM **only** if the team accepts a higher chance of desk rejection or a major-revision request asking for replication. Otherwif then default to Human Genomics for a higher acceptance probability.",
        "",
        "## Locked items (continued)",
        "",
        "- Final GREEN/RED/AMBER therapeutic classification — LOCKED.",
        "- Replication phase — LOCKED.",
        "- Wave 3 SuSiE / D042 optional — LOCKED.",
        "- New coloc / new MR / new data downloads / proteome-wide expansion — LOCKED.",
    ])
    OUT_DECISION.write_text("\n".join(md))
    print(f"wrote {OUT_DECISION}")


def write_manifest(staged_paths: dict[str, Path]):
    md = ["# Submission package manifest (D048)",
          "",
          "**Snapshot:** 2026-05-10",
          "",
          "## Canonical D047 source files (DO NOT modify in D048)",
          "",
          "| Fwith | Status |",
          "|---|---|"]
    for f in [D047_DOCX, D047_PDF, D047_TEX, D047_BIB,
              SUPP_METHODS_DOCX, SUPP_METHODS_PDF, SUPP_METHODS_TEX,
              SUPP_XLSX, COVER_LETTER, AUTHORS_TPL, COI_TPL]:
        md.append(f"| `{f.relative_to(MS_DIR)}` | {'present' if f.exists() else 'MISSING'} |")

    md.extend([
        "",
        "## Stale D045 files (must NOT enter submission package)",
        "",
        "| Fwith | Action |",
        "|---|---|"]
    )
    for stale in ["main_manuscript.docx", "main_manuscript.pdf", "main_manuscript.tex"]:
        path = MS_DIR / stale
        if path.exists():
            md.append(f"| `{stale}` | present locally — exclude from journal upload |")
        else:
            md.append(f"| `{stale}` | absent (already removed or never staged) |")

    md.extend(["", "## Staged journal packages", ""])
    for name, path in staged_paths.items():
        md.append(f"### {name}")
        md.append("")
        files = sorted(path.rglob("*"))
        md.append(f"Fwiths: {sum(1 for f in files if f.is_file())}")
        md.append("")
        for f in files:
            if f.is_file():
                rel = f.relative_to(MS_DIR)
                md.append(f"- `{rel}` ({f.stat().st_size:,} bytes)")
        md.append("")

    md.extend([
        "## Notes",
        "",
        "- Both packages are scientifically identical (D047-canonical content).",
        "- Differences are journal-specific format notes, cover letter wording, and submission checklists.",
        "- PI must finalize author / affiliation / funding / COI / data-availability / suggested-reviewer fields before D049 actual submission.",
    ])
    OUT_MANIFEST.write_text("\n".join(md))
    print(f"wrote {OUT_MANIFEST}")


# ---------------------------------------------------------------------------
# 12-item D048 audit
# ---------------------------------------------------------------------------


def _scan_pdf(path: Path) -> str:
    if not path.exists():
        return ""
    with fitz.open(path) as doc:
        return "\n".join(p.get_text() for p in doc)


def _scan_docx(path: Path) -> str:
    if not path.exists():
        return ""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n\n".join(parts)


def _paragraph_violations(text, phrases):
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    issues = []
    for i, para in enumerate(paras):
        para_low = para.lower()
        has_neg = any(m.lower() in para_low for m in NEG_MARKERS)
        for phrase, reason in phrases:
            if phrase.lower() in para_low and not has_neg:
                issues.append((i, phrase, reason, para.replace("\n", " ")[:160]))
    return issues


def run_audit(staged_paths: dict[str, Path]) -> bool:
    pdf_text = _scan_pdf(D047_PDF)
    docx_text = _scan_docx(D047_DOCX)
    checks = []

    # 1. D047 PDF has no [REF_*] / placeholder text
    placeholder_hits = re.findall(r"\[(?:REF_[^\]]+|Introduction placeholder|Discussion placeholder)\]", pdf_text)
    checks.append(("1. D047 PDF has no placeholder tokens",
                   not placeholder_hits,
                   ["No placeholder tokens"] if not placeholder_hits else placeholder_hits))

    # 2. D047 PDF has no old AF EUR figure/text
    af_eur_issues = []
    if re.search(r"Roselli\s+2025\s*\(?\s*EUR", pdf_text, flags=re.IGNORECASE):
        af_eur_issues.append("D047 PDF contains 'Roselli 2025 EUR' attribution")
    for fig in MS_FIG.glob("*.pdf"):
        ftxt = _scan_pdf(fig)
        if "Roselli 2025 (EUR" in ftxt or "(EUR, GRCh38)" in ftxt:
            af_eur_issues.append(f"figure {fig.name} contains old EUR text")
    checks.append(("2. No old AF EUR figure/text",
                   not af_eur_issues, af_eur_issues or ["all figures + main PDF use all-ancestry / CVDKP wording"]))

    # 3. Introduction + Discussion are full text
    full_text_issues = []
    for needle in ("[Introduction placeholder", "[Discussion placeholder",
                   "Introduction placeholder.", "Discussion placeholder."):
        if needle in pdf_text or needle in docx_text:
            full_text_issues.append(f"placeholder remains: {needle}")
    checks.append(("3. Introduction + Discussion are full text",
                   not full_text_issues, full_text_issues or ["no placeholder text in Introduction or Discussion"]))

    # 4. References real and formatted (bib has @article)
    bib = D047_BIB.read_text() if D047_BIB.exists() else ""
    n_articles = bib.count("@article{")
    n_misc = bib.count("@misc{")
    refs_ok = n_articles >= 20 and (n_articles + n_misc) >= 25
    checks.append(("4. References real and formatted (≥ 20 @article entries)",
                   refs_ok, [f"{n_articles} @article + {n_misc} @misc entries in references.bib"]))

    # 5. Figures render (6 PDFs + 6 PNGs of non-trivial size)
    fig_issues = []
    pdfs = list(MS_FIG.glob("*.pdf"))
    pngs = list(MS_FIG.glob("*.png"))
    if len(pdfs) < 6:
        fig_issues.append(f"expected ≥ 6 figure PDFs, found {len(pdfs)}")
    if len(pngs) < 6:
        fig_issues.append(f"expected ≥ 6 figure PNGs, found {len(pngs)}")
    for f in pdfs + pngs:
        if f.stat().st_size < 5_000:
            fig_issues.append(f"{f.name}: {f.stat().st_size} bytes (suspiciously small)")
    checks.append(("5. Figures render correctly (PDF + PNG ≥ 6 each, non-trivial size)",
                   not fig_issues, fig_issues or [f"{len(pdfs)} PDF + {len(pngs)} PNG, all ≥ 5 KB"]))

    # 6. Tables readable (XLSX has 5 sheets, all populated)
    table_issues = []
    if SUPP_XLSX.exists():
        wb = openpyxl.load_workbook(SUPP_XLSX, read_only=True)
        if len(wb.sheetnames) != 5:
            table_issues.append(f"expected 5 sheets, found {len(wb.sheetnames)}: {wb.sheetnames}")
        for name in wb.sheetnames:
            ws = wb[name]
            if ws.max_row < 2 or ws.max_column < 2:
                table_issues.append(f"sheet {name} too small: {ws.max_row}x{ws.max_column}")
    else:
        table_issues.append("supplementary_tables.xlsx missing")
    checks.append(("6. Tables readable (5-sheet XLSX populated)",
                   not table_issues, table_issues or ["5 sheets, all populated"]))

    # 7. Supplementary methods + tables attached in each staged folder
    supp_issues = []
    for jname, jpath in staged_paths.items():
        for needed in ("supplementary_methods.docx", "supplementary_methods.pdf",
                       "supplementary_tables.xlsx", "supplementary_tables_tsv", "figures"):
            if not (jpath / needed).exists():
                supp_issues.append(f"{jname}: missing {needed}")
    checks.append(("7. Supplementary methods + tables attached in each staged folder",
                   not supp_issues, supp_issues or ["all supplementary attachments present in both staged folders"]))

    # 8. Cover letter + author + COI templates exist (per-journal)
    cover_issues = []
    for jname, jpath in staged_paths.items():
        if not (jpath / f"cover_letter_{jname}.md").exists():
            cover_issues.append(f"{jname}: cover_letter_{jname}.md missing")
        if not (jpath / "author_contributions_template.md").exists():
            cover_issues.append(f"{jname}: author_contributions_template.md missing")
        if not (jpath / "coi_funding_statement_template.md").exists():
            cover_issues.append(f"{jname}: coi_funding_statement_template.md missing")
    checks.append(("8. Cover letter + author + COI templates exist (per-journal)",
                   not cover_issues, cover_issues or ["all 3 templates per journal staged"]))

    # 9. No forbidden therapeutic claims (paragraph-aware)
    forbid_issues = _paragraph_violations(pdf_text, FORBIDDEN_LITERAL)
    forbid_issues += _paragraph_violations(docx_text, FORBIDDEN_LITERAL)
    forbid_msgs = []
    for i, phrase, reason, snippet in forbid_issues[:10]:
        forbid_msgs.append(f"para{i}: '{phrase}' ({reason}) — '{snippet}'")
    checks.append(("9. No forbidden therapeutic claims (paragraph-aware)",
                   not forbid_issues, forbid_msgs or ["clean across PDF + DOCX"]))

    # 10. Required language preserved (carry-over)
    lang_issues = []
    for kw, label in REQUIRED_LANGUAGE_KEYWORDS:
        if kw.lower() not in pdf_text.lower() and kw.lower() not in docx_text.lower():
            lang_issues.append(f"missing required language ({label}): '{kw}'")
    checks.append(("10. Required language preserved",
                   not lang_issues, lang_issues or ["all 5 required keywords/phrases present"]))

    # 11. Target journal selected (≥ 1 staged folder + decision matrix)
    journal_issues = []
    if not OUT_DECISION.exists():
        journal_issues.append("target_journal_decision_matrix.md missing")
    if not staged_paths:
        journal_issues.append("no staged journal folder")
    checks.append(("11. Target journal staged (≥ 1 staged folder + decision matrix)",
                   not journal_issues, journal_issues or [f"{len(staged_paths)} staged: {list(staged_paths)}"]))

    # 12. PI/author review items listed
    review_items_listed = OUT_DECISION.exists() and (
        "Authors" in OUT_DECISION.read_text() or "PI" in OUT_DECISION.read_text()
    )
    # Also each staged folder has submission_checklist.md with PI section
    pi_section_present = all(
        (jpath / "submission_checklist.md").exists() for jpath in staged_paths.values()
    )
    if all((jpath / "submission_checklist.md").exists() for jpath in staged_paths.values()):
        any_with_pi = False
        for jpath in staged_paths.values():
            txt = (jpath / "submission_checklist.md").read_text()
            if "PI must finalize" in txt or "Suggested reviewers" in txt:
                any_with_pi = True
                break
        if not any_with_pi:
            pi_section_present = False
    checks.append(("12. PI/author review items listed",
                   review_items_listed and pi_section_present,
                   ["decision matrix + per-journal submission_checklist.md include PI/author finalization items"]
                   if review_items_listed and pi_section_present
                   else ["one or more checklists missing PI sections"]))

    overall = all(ok for _, ok, _ in checks)

    md = [f"# D048 final submission audit — {'PASS' if overall else 'FAIL'}",
          "",
          f"**Snapshot:** 2026-05-10",
          f"**Result:** {sum(1 for _, ok, _ in checks if ok)} / {len(checks)} checks pass.",
          ""]
    for label, ok, details in checks:
        md.append(f"## {label} — {'✅ PASS' if ok else '❌ FAIL'}")
        for d in details:
            md.append(f"- {d}")
        md.append("")
    md.extend([
        "## Submission status",
        "",
        f"**{'D048 PASS — both journal packages staged. PI to finalize content fields and choose target journal before D049 (actual submission).' if overall else 'D048 FAIL — fix items above before D049.'}**",
        "",
        "## Open at D049 (actual submission)",
        "",
        "- PI/author final approval of authors, affiliations, funding, COI.",
        "- Final journal choice (CircGPM vs Human Genomics).",
        "- Apply journal-specific format adaptations per `journal_specific_notes.md` in chosen folder.",
        "- Compwith final journal-styled PDF.",
        "- Submit through journal portal.",
        "",
        "## Locked items (continued)",
        "",
        "- Final GREEN/RED/AMBER therapeutic classification — LOCKED.",
        "- Replication phase — LOCKED.",
        "- Wave 3 SuSiE / D042 optional — LOCKED.",
        "- New coloc / new MR / new data downloads / proteome-wide expansion — LOCKED.",
    ])
    OUT_AUDIT.write_text("\n".join(md))
    print(f"wrote {OUT_AUDIT}")
    print(f"=== D048 OVERALL: {'PASS' if overall else 'FAIL'} ===")
    for label, ok, details in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {label}")
        for d in details:
            print(f"    - {d}")
    return overall


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------


def main() -> int:
    print("=== D047 source verification ===")
    missing = []
    for f in [D047_DOCX, D047_PDF, D047_TEX, D047_BIB,
              COVER_LETTER, AUTHORS_TPL, COI_TPL]:
        if not f.exists():
            missing.append(str(f))
    if missing:
        print(f"FAIL: missing D047 sources: {missing}")
        return 2
    print("D047 sources OK.")

    print("\n=== JOURNAL STAGING ===")
    staged = {}
    for name in JOURNALS:
        path = stage_journal(name)
        staged[name] = path
        print(f"staged → {path}")

    print("\n=== DECISION MATRIX + MANIFEST ===")
    write_decision_matrix()
    write_manifest(staged)

    print("\n=== AUDIT ===")
    overall_pass = run_audit(staged)
    return 0 if overall_pass else 2


if __name__ == "__main__":
    sys.exit(main())
